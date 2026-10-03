import io
import logging

logger = logging.getLogger("skillgap.parser")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from PDF file bytes using pypdf."""
    text_content = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    except Exception as e:
        logger.warning(f"Error parsing PDF with pypdf: {e}")
        # Secondary fallback using pdfplumber if installed
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text()
                    if txt:
                        text_content.append(txt)
        except Exception as e2:
            logger.error(f"Failed all PDF extractors: {e2}")

    return "\n".join(text_content).strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from DOCX file bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    full_text.append(row_text)
        return "\n".join(full_text).strip()
    except Exception as e:
        logger.error(f"Error parsing DOCX file: {e}")
        return ""

def parse_document(file_bytes: bytes, filename: str) -> str:
    """Detect file type and extract text content."""
    filename_lower = filename.lower()
    
    if filename_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    else:
        # Default try utf-8 text decoding
        try:
            return file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            return ""
